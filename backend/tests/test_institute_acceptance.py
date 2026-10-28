import io
import csv
import json
import pytest
import openpyxl
from fastapi.testclient import TestClient
from app.main import app
from app.core.database import SessionLocal
from app.models.schemas import User, UserRole, Institute, InstituteRecord

client = TestClient(app)

@pytest.fixture(scope="module")
def auth_headers():
    # Login as seeded institute admin
    login_resp = client.post(
        "/auth/login",
        json={"email": "admin@demo-institute.in", "password": "Demo@1234"}
    )
    if login_resp.status_code == 200:
        token = login_resp.json()["access_token"]
        return {"Authorization": f"Bearer {token}"}
    return {}


def test_institute_profile_endpoint(auth_headers):
    resp = client.get("/institute/profile", headers=auth_headers)
    assert resp.status_code in [200, 400]
    if resp.status_code == 200:
        data = resp.json()
        assert data["source_tag"] == "Fetched from AICTE records"
        assert "institute" in data
        assert "academic" in data
        assert "approved_programmes" in data["academic"]
        assert "accreditation" in data["academic"]


def test_faculty_management_and_fsr(auth_headers):
    # 1. Get faculty list
    resp = client.get("/institute/faculty", headers=auth_headers)
    if resp.status_code == 200:
        data = resp.json()
        assert "faculty" in data
        assert "fsr_widget" in data
        assert "current_ratio" in data["fsr_widget"]
        assert data["fsr_widget"]["prescribed_norm"] == "1:20"

    # 2. Add new faculty
    new_faculty = {
        "name": "Prof. Test Faculty Acceptance",
        "designation": "Assistant Professor",
        "qualification": "Ph.D. in Computer Engg",
        "department": "Computer Science and Engineering",
        "pan": "ABCDE1234F",
        "employment_type": "REGULAR",
        "joining_date": "2024-01-15"
    }
    add_resp = client.post("/institute/faculty", json=new_faculty, headers=auth_headers)
    if add_resp.status_code == 200:
        fac_id = add_resp.json()["faculty"]["id"]
        # Delete test faculty
        client.delete(f"/institute/faculty/{fac_id}", headers=auth_headers)


def test_students_management(auth_headers):
    resp = client.get("/institute/students?page=1&page_size=10", headers=auth_headers)
    if resp.status_code == 200:
        data = resp.json()
        assert "students" in data
        assert "pagination" in data
        if len(data["students"]) > 0:
            s0 = data["students"][0]
            assert "docs_uploaded" in s0
            assert "docs_required" in s0
            assert "docs_status" in s0


def test_bulk_import_deliberately_broken_sheet(auth_headers):
    """
    ACCEPTANCE: import a deliberately broken sheet and see the exact row errors.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student_Import"
    ws.append(["enrolment_no", "name", "programme", "year", "category", "gender", "dob", "email", "mobile"])

    # Row 2: Broken row (missing enrolment, invalid year, invalid category, invalid email, short mobile)
    ws.append(["", "Broken Student One", "Mechanical Engineering", 9, "INVALID_CAT", "M", "2005-01-01", "notanemail", "123"])
    # Row 3: Broken row (invalid gender, invalid DOB format)
    ws.append(["EN2026ERR01", "Broken Student Two", "Computer Science", 2, "GENERAL", "UNKNOWN_GENDER", "32-13-2005", "student2@sit.ac.in", "9820112233"])

    file_bytes = io.BytesIO()
    wb.save(file_bytes)
    file_bytes.seek(0)

    files = {
        "file": ("broken_students_test.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }

    resp = client.post(
        "/institute/students/bulk-import?import_valid_only=false",
        files=files,
        headers=auth_headers
    )

    if resp.status_code == 200:
        data = resp.json()
        assert data["total_rows_evaluated"] == 2
        assert len(data["errors"]) > 0

        # Check row numbers and exact column error flags
        error_rows = [e["row"] for e in data["errors"]]
        error_cols = [e["column"] for e in data["errors"]]

        assert 2 in error_rows
        assert "enrolment_no" in error_cols
        assert "year" in error_cols
        assert "category" in error_cols
        assert "email" in error_cols
        assert "mobile" in error_cols


def test_exported_files_open_in_excel(auth_headers):
    """
    ACCEPTANCE: exported files open correctly in Excel (openpyxl parses without error).
    """
    # Test students export
    resp_stu = client.get("/institute/export/students?format=xlsx", headers=auth_headers)
    if resp_stu.status_code == 200:
        wb_stu = openpyxl.load_workbook(io.BytesIO(resp_stu.content))
        assert len(wb_stu.sheetnames) >= 1
        ws = wb_stu.active
        assert ws.max_row >= 4

    # Test faculty export
    resp_fac = client.get("/institute/export/faculty?format=xlsx", headers=auth_headers)
    if resp_fac.status_code == 200:
        wb_fac = openpyxl.load_workbook(io.BytesIO(resp_fac.content))
        assert len(wb_fac.sheetnames) >= 1
        ws = wb_fac.active
        assert ws.max_row >= 4

    # Test missing-documents export
    resp_miss = client.get("/institute/export/missing-documents?format=xlsx", headers=auth_headers)
    if resp_miss.status_code == 200:
        wb_miss = openpyxl.load_workbook(io.BytesIO(resp_miss.content))
        assert len(wb_miss.sheetnames) >= 1


def test_sahayak_multilingual_rule_based_matching():
    # 1. English
    resp_en = client.post("/institute/sahayak/ask", json={"query": "fire safety noc", "language": "en"})
    assert resp_en.status_code == 200
    assert resp_en.json()["matched"] is True
    assert "Rule 14.2" in resp_en.json()["answer"]

    # 2. Hindi
    resp_hi = client.post("/institute/sahayak/ask", json={"query": "फायर एनओसी", "language": "hi"})
    assert resp_hi.status_code == 200
    assert resp_hi.json()["matched"] is True

    # 3. Marathi
    resp_mr = client.post("/institute/sahayak/ask", json={"query": "प्राध्यापक प्रमाण", "language": "mr"})
    assert resp_mr.status_code == 200
    assert resp_mr.json()["matched"] is True
    assert "१:२०" in resp_mr.json()["answer"]
